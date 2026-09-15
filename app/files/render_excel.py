"""Excel 渲染器:openpyxl,大标题/表头样式/冻结窗格/列宽自适应/预算合计公式/打印设置。"""
from __future__ import annotations

import re
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.properties import PageSetupProperties

from app.models.intermediates import ExcelDocIR, SheetIR
from app.utils.text import cjk_width

PRIMARY_FILL = PatternFill("solid", fgColor="1F4E79")
ALT_FILL = PatternFill("solid", fgColor="F2F6FA")
TOTAL_FILL = PatternFill("solid", fgColor="FFF2CC")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HEADER_FONT = Font(name="微软雅黑", size=11, bold=True, color="FFFFFF")
BODY_FONT = Font(name="微软雅黑", size=11, color="333333")
TITLE_FONT = Font(name="黑体", size=14, bold=True, color="1F4E79")

_NUMERIC_RE = re.compile(r"^[\d,，.%\s~\-+～－]+$")


def render(ir: ExcelDocIR, out_path: Path) -> None:
    wb = Workbook()
    for i, sheet_ir in enumerate(ir.sheets[:3]):
        ws = wb.active if i == 0 else wb.create_sheet()
        ws.title = sheet_ir.name[:12]
        if sheet_ir.kind == "notes":
            _render_notes(ws, sheet_ir)
        else:
            _render_sheet(ws, sheet_ir)
    wb.save(str(out_path))


def _render_sheet(ws, sheet_ir: SheetIR) -> None:
    headers = [h for h in sheet_ir.headers[:8]]
    rows = [row[: len(headers) or 8] for row in sheet_ir.rows[:40]]
    ncols = max(len(headers), 1)
    last_col = get_column_letter(ncols)

    # 大标题
    ws.merge_cells(f"A1:{last_col}1")
    cell = ws["A1"]
    cell.value = sheet_ir.title_row or sheet_ir.name
    cell.font = TITLE_FONT
    cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 26

    # 表头
    for c, text in enumerate(headers, start=1):
        cell = ws.cell(row=2, column=c, value=text)
        cell.font = HEADER_FONT
        cell.fill = PRIMARY_FILL
        cell.border = BORDER
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[2].height = 22

    # 数据行
    numeric_cols: set[int] = set()
    for r, row in enumerate(rows, start=3):
        for c, value in enumerate(row, start=1):
            text = str(value)[:30] if value is not None else ""
            cell = ws.cell(row=r, column=c, value=text)
            cell.font = BODY_FONT
            cell.border = BORDER
            if r % 2 == 0:
                cell.fill = ALT_FILL
            if _NUMERIC_RE.match(text) and text:
                cell.alignment = Alignment(horizontal="right", vertical="center")
                numeric_cols.add(c)
            else:
                cell.alignment = Alignment(vertical="center", wrap_text=True)

    # 预算表:数值列加 SUM 合计行
    if sheet_ir.kind == "budget" and numeric_cols and rows:
        total_row = 3 + len(rows)
        ws.cell(row=total_row, column=1, value="合计")
        for c in sorted(numeric_cols):
            ws.cell(row=total_row, column=c,
                    value=f"=SUM({get_column_letter(c)}3:{get_column_letter(c)}{total_row - 1})")
        for c in range(1, ncols + 1):
            cell = ws.cell(row=total_row, column=c)
            cell.font = Font(name="微软雅黑", size=11, bold=True)
            cell.fill = TOTAL_FILL
            cell.border = BORDER

    # 列宽(按 CJK 宽度估算)
    for c in range(1, ncols + 1):
        width = 8
        for r, row in enumerate(rows, start=3):
            value = row[c - 1] if c - 1 < len(row) else ""
            width = max(width, min(cjk_width(str(value)) + 2, 50))
        for h in headers:
            width = max(width, min(cjk_width(h) + 2, 50))
        ws.column_dimensions[get_column_letter(c)].width = width

    ws.freeze_panes = "A3"
    ws.sheet_view.showGridLines = False
    _print_setup(ws)


def _render_notes(ws, sheet_ir: SheetIR) -> None:
    """说明页:每行一条文本,跨列合并、自动换行。"""
    lines = sheet_ir.rows[:40]
    ncols = max(len(sheet_ir.headers or [1]), 6)
    last_col = get_column_letter(ncols)

    ws.merge_cells(f"A1:{last_col}1")
    cell = ws["A1"]
    cell.value = sheet_ir.title_row or sheet_ir.name
    cell.font = TITLE_FONT
    cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 26

    r = 2
    for row in lines:
        texts = [t for t in row if str(t).strip()]
        if not texts:
            continue
        line = " ".join(str(t) for t in texts)[:300]
        ws.merge_cells(f"A{r}:{last_col}{r}")
        cell = ws.cell(row=r, column=1, value=line)
        cell.font = BODY_FONT
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        ws.row_dimensions[r].height = max(18, 16 * (1 + len(line) // 60))
        r += 1

    ws.column_dimensions["A"].width = 100
    ws.sheet_view.showGridLines = False
    _print_setup(ws)


def _print_setup(ws) -> None:
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
